import io
from types import SimpleNamespace

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from polish_pptx.core import CONSISTENCY_PROMPT, count_diagram_shapes, extract_svg, polish_presentation

SVG = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 50"><rect width="90" height="40" fill="red"/></svg>'


def make_prs():
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    for i in range(3):
        s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(i), Inches(1), Inches(1), Inches(1))
    prs.slides.add_slide(prs.slide_layouts[6]).shapes.add_textbox(0, 0, Inches(1), Inches(1))
    return prs


def test_count():
    prs = make_prs()
    assert [count_diagram_shapes(s) for s in prs.slides] == [3, 0]


def test_extract_svg():
    assert extract_svg("```svg\n" + SVG + "\n```") == SVG


def test_polish():
    prs = make_prs()

    class Msgs:
        def create(self, **kw):
            return SimpleNamespace(content=[SimpleNamespace(type="text", text=SVG)])

    client = SimpleNamespace(messages=Msgs())
    imgs = [Image.new("RGB", (100, 50))] * 2
    assert polish_presentation(prs, client, images=imgs, log=lambda m: None) == 1
    assert len(prs.slides[0].shapes) == 4
    assert len(prs.slides[1].shapes) == 1


def test_polish_selected_slides():
    prs = make_prs()

    class Msgs:
        def create(self, **kw):
            return SimpleNamespace(content=[SimpleNamespace(type="text", text=SVG)])

    client = SimpleNamespace(messages=Msgs())
    imgs = [Image.new("RGB", (100, 50))] * 2
    assert polish_presentation(prs, client, images=imgs, slides=[2], log=lambda m: None) == 1
    assert len(prs.slides[0].shapes) == 3
    assert len(prs.slides[1].shapes) == 2


def test_polish_sends_earlier_slides_as_examples():
    prs = make_prs()
    calls = []

    class Msgs:
        def create(self, **kw):
            calls.append(kw["messages"])
            return SimpleNamespace(content=[SimpleNamespace(type="text", text=SVG)])

    client = SimpleNamespace(messages=Msgs())
    imgs = [Image.new("RGB", (100, 50))] * 2
    assert polish_presentation(prs, client, images=imgs, slides=[1, 2], log=lambda m: None) == 2
    assert [m["role"] for m in calls[0]] == ["user"]
    assert [m["role"] for m in calls[1]] == ["user", "assistant", "user"]
    assert calls[1][1]["content"][0]["text"] == SVG
    assert calls[1][2]["content"][-1]["text"].endswith(CONSISTENCY_PROMPT)
    assert CONSISTENCY_PROMPT not in calls[0][0]["content"][-1]["text"]


def test_request_svg_limits_examples():
    from polish_pptx.core import MAX_EXAMPLES, request_svg

    calls = []

    class Msgs:
        def create(self, **kw):
            calls.append(kw["messages"])
            return SimpleNamespace(content=[SimpleNamespace(type="text", text=SVG)])

    client = SimpleNamespace(messages=Msgs())
    history = []
    for _ in range(MAX_EXAMPLES + 2):
        request_svg(client, Image.new("RGB", (100, 50)), history=history)
    assert len(history) == MAX_EXAMPLES + 2
    assert len(calls[-1]) == 2 * MAX_EXAMPLES + 1
