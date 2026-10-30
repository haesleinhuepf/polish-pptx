import io
import re

from pptx.enum.shapes import MSO_SHAPE_TYPE

DEFAULT_MODEL = "claude-sonnet-5"

PROMPT = (
    "This is a screenshot of a presentation slide. Polish the figure on it: focus on the "
    "arrows, boxes, circles and similar diagram shapes (and their labels) and redraw them "
    "cleanly and consistently aligned. Ignore the headline/title text and layout logos. "
    "Return only the polished figure as a single SVG document (no explanation, no markdown), "
    "with a viewBox of the same aspect ratio as the screenshot and a transparent background."
)


def is_diagram_shape(shape):
    """True for autoshapes, connectors and freeforms (not text boxes, pictures, placeholders)."""
    if shape.is_placeholder:
        return False
    return shape.shape_type in (
        MSO_SHAPE_TYPE.AUTO_SHAPE,
        MSO_SHAPE_TYPE.FREEFORM,
        MSO_SHAPE_TYPE.LINE,
    ) or shape.element.tag.endswith("}cxnSp")


def count_diagram_shapes(slide):
    return sum(1 for s in slide.shapes if is_diagram_shape(s))


def image_to_base64_png(image):
    import base64

    buf = io.BytesIO()
    image.convert("RGB").save(buf, format="PNG")
    return base64.standard_b64encode(buf.getvalue()).decode()


def extract_svg(text):
    m = re.search(r"<svg\b.*?</svg>", text, re.DOTALL | re.IGNORECASE)
    if not m:
        raise ValueError("Model response did not contain an SVG")
    return m.group(0)


def request_svg(client, image, model=DEFAULT_MODEL):
    response = client.messages.create(
        model=model,
        max_tokens=16000,
        messages=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "image",
                        "source": {
                            "type": "base64",
                            "media_type": "image/png",
                            "data": image_to_base64_png(image),
                        },
                    },
                    {"type": "text", "text": PROMPT},
                ],
            }
        ],
    )
    text = "".join(b.text for b in response.content if getattr(b, "type", "") == "text")
    return extract_svg(text)


def svg_to_png_bytes(svg, width_px=1920):
    import cairosvg

    return cairosvg.svg2png(bytestring=svg.encode("utf-8"), output_width=width_px)


def add_png_to_slide(prs, slide, png_bytes):
    """Add the PNG to the slide, fitted and centered within the slide."""
    from PIL import Image

    w, h = Image.open(io.BytesIO(png_bytes)).size
    scale = min(prs.slide_width / w, prs.slide_height / h)
    pw, ph = int(w * scale), int(h * scale)
    left = (prs.slide_width - pw) // 2
    top = (prs.slide_height - ph) // 2
    return slide.shapes.add_picture(io.BytesIO(png_bytes), left, top, pw, ph)


def polish_presentation(prs, client, images=None, model=DEFAULT_MODEL, min_shapes=3, log=print):
    """Polish all slides having at least ``min_shapes`` diagram shapes. Returns polished count."""
    targets = [
        i for i, s in enumerate(prs.slides) if count_diagram_shapes(s) >= min_shapes
    ]
    if not targets:
        return 0
    if images is None:
        from render_pptx import render_slides

        images = render_slides(prs)
    done = 0
    for i in targets:
        slide = prs.slides[i]
        log(f"Polishing slide {i + 1}")
        try:
            svg = request_svg(client, images[i], model)
            png = svg_to_png_bytes(svg)
            add_png_to_slide(prs, slide, png)
            done += 1
        except Exception as e:  # keep going with other slides
            log(f"  slide {i + 1} skipped: {e}")
    return done
